import os
import sys

sys.stdout.reconfigure(encoding='utf-8')
from pptx import Presentation

pptx_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "..", "FixNear_BaoCao_ASM_7Buoc_v2.pptx")
if not os.path.exists(pptx_path):
    pptx_path = "FixNear_BaoCao_ASM_7Buoc_v2.pptx"

prs = Presentation(pptx_path)
print(f"Total slides: {len(prs.slides)}")
for idx, slide in enumerate(prs.slides):
    title = slide.shapes.title.text if slide.shapes.title else "No Title"
    texts = []
    for shape in slide.shapes:
        if shape.has_text_frame and shape != slide.shapes.title:
            texts.append(shape.text.strip().replace("\n", " "))
    preview = " | ".join(texts)[:120]
    print(f"Slide {idx+1}: [{title}] -> {preview}")
