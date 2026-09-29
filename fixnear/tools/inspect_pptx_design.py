import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

pptx_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "..", "FixNear_BaoCao_ASM_7Buoc_v2.pptx")
prs = Presentation(pptx_path)

print(f"Slide width: {prs.slide_width.inches} inches, height: {prs.slide_height.inches} inches")
for idx in range(min(5, len(prs.slides))):
    slide = prs.slides[idx]
    print(f"\n--- SLIDE {idx+1} ---")
    for s in slide.shapes:
        if s.has_text_frame:
            for p in s.text_frame.paragraphs:
                if p.text.strip():
                    font_name = p.font.name if p.font else None
                    font_size = p.font.size.pt if p.font and p.font.size else None
                    print(f"  [Text]: {p.text.strip()[:60]} (Font: {font_name}, Size: {font_size})")
