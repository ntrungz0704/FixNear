import os
import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation

pptx_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "..", "FixNear_BaoCao_ASM_7Buoc_v2.pptx")
prs = Presentation(pptx_path)

slide16 = prs.slides[15]
print("--- SLIDE 16 DETAILS ---")
for s in slide16.shapes:
    print(f"Shape: {s.name} | Type: {s.shape_type} | L={s.left.inches:.2f}, T={s.top.inches:.2f}, W={s.width.inches:.2f}, H={s.height.inches:.2f}")
    if s.has_text_frame:
        for p in s.text_frame.paragraphs:
            if p.text.strip():
                print(f"   P: {p.text.strip()[:80]}")
